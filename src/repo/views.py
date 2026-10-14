import codecs
import csv
import io
from typing import TYPE_CHECKING, Any

if TYPE_CHECKING:
    from curation.models import Curation

from django.db.models import Prefetch, QuerySet
from django.http import HttpRequest, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, render
from django.utils import timezone
from django.views.generic import DetailView
from django_tables2 import RequestConfig
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from common.history import resolve_changes
from common.tables import HistoryTable
from common.views import SearchListView
from curation.constants.models.common import Status
from curation.models import Evidence
from repo.models import PublishedCuration
from repo.serializers import (
    flatten_classification_rows,
    flatten_evidence,
    flatten_published_curation,
    serialize_published_curation,
)
from repo.tables import PublishedCurationTable


def is_superseded(published_curation: PublishedCuration) -> bool:
    """Returns True if a direct or transitive copy is also published."""

    def _has_published_copy(curation: "Curation") -> bool:
        for copy in curation.copies.all():  # type: ignore
            if copy.status == Status.PUBLISHED:
                return True
            if _has_published_copy(copy):
                return True
        return False

    return _has_published_copy(published_curation.curation)


def get_superseding(published_curation: PublishedCuration) -> PublishedCuration | None:
    """Returns the most recently published descendant, or None if not superseded."""

    def _find_latest(curation: "Curation") -> PublishedCuration | None:
        result = None
        for copy in curation.copies.all():  # type: ignore
            if copy.status == Status.PUBLISHED:
                candidate = copy.publication
                if result is None or candidate.published_at > result.published_at:
                    result = candidate
            deeper = _find_latest(copy)
            if deeper is not None and (
                result is None or deeper.published_at > result.published_at
            ):
                result = deeper
        return result

    return _find_latest(published_curation.curation)


class PublishedCurationList(SearchListView):
    model = PublishedCuration
    template_name = "repo/list.html"
    ordering = ["-published_at"]
    table_class = PublishedCurationTable
    search_fields = [
        "curation__slug",
        "curation__allele__name",
        "curation__haplotype__name",
        "curation__disease__name",
    ]
    table_pagination = {"per_page": 25}


class PublishedCurationDetail(DetailView):
    model = PublishedCuration
    template_name = "repo/detail.html"

    def get_object(self, queryset: QuerySet[Any] | None = None) -> PublishedCuration:  # ruff: ignore[unused-method-argument]
        """Get PublishedCuration by the curation's slug.

        Returns:
            PublishedCuration instance matching the slug from URL kwargs.
        """
        curation_slug = self.kwargs.get("curation_slug")
        return get_object_or_404(
            PublishedCuration.objects.select_related(
                "curation__allele", "curation__haplotype", "curation__disease"
            ).prefetch_related("curation__haplotype__alleles"),
            curation__slug=curation_slug,
        )

    def get_context_data(self, **kwargs) -> dict[str, Any]:
        """Add curation to context for template convenience.

        Returns:
            Context dict with curation, superseded, and superseding keys added.
        """
        context = super().get_context_data(**kwargs)
        context["curation"] = self.object.curation
        context["superseded"] = is_superseded(self.object)
        context["superseding"] = get_superseding(self.object)
        return context


class PublishedCurationHistory(DetailView):
    model = PublishedCuration
    template_name = "repo/history.html"

    def get_object(self, queryset: QuerySet[Any] | None = None) -> PublishedCuration:  # ruff: ignore[unused-method-argument]
        return get_object_or_404(
            PublishedCuration, curation__slug=self.kwargs["curation_slug"]
        )

    def get_context_data(self, **kwargs) -> dict[str, Any]:
        context = super().get_context_data(**kwargs)
        history_table = HistoryTable(
            self.object.history.all(),
            change_url_name="repo-change",
            change_url_slug1=self.object.curation.slug,
            hide_user=True,
        )
        RequestConfig(self.request).configure(history_table)
        context["history_table"] = history_table
        context["curation"] = self.object.curation
        return context


class PublishedCurationChange(DetailView):
    model = PublishedCuration
    template_name = "repo/change.html"

    def get_object(self, queryset: QuerySet[Any] | None = None) -> PublishedCuration:  # ruff: ignore[unused-method-argument]
        return get_object_or_404(
            PublishedCuration, curation__slug=self.kwargs["curation_slug"]
        )

    def get_context_data(self, **kwargs) -> dict[str, Any]:
        context = super().get_context_data(**kwargs)
        record = self.object.history.get(history_id=self.kwargs["history_id"])
        prev_record = record.prev_record
        context["record"] = record
        # HLArepo is public, so it doesn't show who changed or published a curation.
        context["changes"] = resolve_changes(
            PublishedCuration, record, prev_record, exclude=["published_by"]
        )
        context["hide_user"] = True
        context["curation"] = self.object.curation
        return context


def _json_attachment(
    published_curations: list[PublishedCuration], filename: str
) -> HttpResponse:
    """Returns curations as an indented, downloadable JSON file.

    The single and bulk downloads share this shape, so one reader handles both.

    Returns:
        JSON response with the export date, the count, and the curations.
    """
    data = {
        "export_date": timezone.now().strftime("%Y-%m-%d"),
        "total_count": len(published_curations),
        "curations": [
            serialize_published_curation(published) for published in published_curations
        ],
    }
    # Indent the file and keep non-ASCII text, since people read it as is.
    response = JsonResponse(
        data, json_dumps_params={"indent": 2, "ensure_ascii": False}
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def download_all_json(request: HttpRequest) -> HttpResponse:
    """Downloads all published curations as JSON.

    Returns:
        JSON response with all published curations and metadata.
    """
    timestamp = timezone.now().strftime("%Y-%m-%d")
    return _json_attachment(
        list(PublishedCuration.objects.order_by("-published_at")),
        filename=f"hla_curations_all_{timestamp}.json",
    )


def _published_for_export() -> QuerySet[PublishedCuration]:
    """Returns every published curation with what the flat files need.

    Returns:
        Published curations, most recently published first, with their related
        rows, evidence, publications, and demographics prefetched so the query
        count doesn't grow with the number of rows.
    """
    return (
        PublishedCuration.objects.select_related(
            "curation__allele",
            "curation__haplotype",
            "curation__disease",
            "curation__copied_from",
        )
        .prefetch_related(
            Prefetch(
                "curation__evidence",
                queryset=Evidence.objects.select_related(
                    "publication"
                ).prefetch_related("demographics"),
            )
        )
        .order_by("-published_at")
    )


def download_all_csv(request: HttpRequest) -> HttpResponse:
    """Downloads all published classifications as CSV, one row per evidence.

    The file starts with a UTF-8 byte order mark so Excel reads non-ASCII text
    correctly.

    Returns:
        CSV response with all published classifications, most recently published
        first.
    """
    published_curations = _published_for_export()
    timestamp = timezone.now().strftime("%Y-%m-%d")
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = (
        f'attachment; filename="hla_curations_all_{timestamp}.csv"'
    )
    response.write(codecs.BOM_UTF8)
    writer = None
    for published in published_curations:
        url = request.build_absolute_uri(published.get_absolute_url())
        for row in flatten_classification_rows(published, url=url):
            if writer is None:
                writer = csv.DictWriter(response, fieldnames=list(row))
                writer.writeheader()
            writer.writerow(row)
    return response


XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _write_sheet(sheet: Worksheet, rows: list[dict[str, object]]) -> None:
    """Writes rows to a sheet under a bold, frozen header row."""
    if not rows:
        return
    sheet.append(list(rows[0]))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    for row in rows:
        sheet.append(list(row.values()))


def download_all_xlsx(request: HttpRequest) -> HttpResponse:
    """Downloads all published curations as an Excel workbook.

    The Curations sheet has one row per curation. The Evidence sheet has one row
    per evidence record, starting with its curation's ID.

    Returns:
        XLSX response with all published curations, most recently published
        first.
    """
    curation_rows = []
    evidence_rows = []
    for published in _published_for_export():
        url = request.build_absolute_uri(published.get_absolute_url())
        curation_rows.append(flatten_published_curation(published, url=url))
        curation_id = {"Curation ID": published.curation.slug}
        evidence_rows.extend(
            {**curation_id, **flatten_evidence(evidence)}
            for evidence in published.curation.evidence.all()
        )

    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])  # Drop the default empty sheet.
    _write_sheet(workbook.create_sheet("Curations"), curation_rows)
    _write_sheet(workbook.create_sheet("Evidence"), evidence_rows)
    buffer = io.BytesIO()
    workbook.save(buffer)

    timestamp = timezone.now().strftime("%Y-%m-%d")
    response = HttpResponse(buffer.getvalue(), content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = (
        f'attachment; filename="hla_curations_all_{timestamp}.xlsx"'
    )
    return response


def download_single_json(request: HttpRequest, curation_slug: str) -> HttpResponse:
    """Downloads a single published curation as JSON.

    Returns:
        JSON response with the specified published curation.
    """
    published = get_object_or_404(PublishedCuration, curation__slug=curation_slug)
    return _json_attachment([published], filename=f"hla_curation_{curation_slug}.json")


def about(request: HttpRequest) -> HttpResponse:
    """Returns the HLArepo about page."""
    return render(request, "repo/about.html")


def contact(request: HttpRequest) -> HttpResponse:
    """Returns the HLArepo contact page."""
    return render(request, "repo/contact.html")


def help_(request: HttpRequest) -> HttpResponse:
    """Returns the HLArepo help page."""
    return render(request, "repo/help.html")


def citing(request: HttpRequest) -> HttpResponse:
    """Returns the HLArepo citing page."""
    return render(request, "repo/citing.html")


def acknowledgements(request: HttpRequest) -> HttpResponse:
    """Returns the HLArepo acknowledgements page."""
    return render(request, "repo/acknowledgements.html")


def collaborators(request: HttpRequest) -> HttpResponse:
    """Returns the HLArepo collaborators page."""
    return render(request, "repo/collaborators.html")
