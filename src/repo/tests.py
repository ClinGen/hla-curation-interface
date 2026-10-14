import codecs
import csv
import io
import json
from datetime import date, datetime, timedelta
from decimal import Decimal
from typing import TYPE_CHECKING, override

from bs4 import BeautifulSoup
from django.conf import settings
from django.contrib.auth.models import User
from django.db import IntegrityError
from django.templatetags.static import static
from django.test import Client, TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from allele.models import Allele
from auth_.models import UserProfile
from common.tests import OpenViewTestMixin, ProtectedViewTestMixin
from curation.constants.models.common import Status
from curation.constants.models.curation import (
    HLA_CURATION_TASKFORCE_ID,
    TASKFORCE_CHOICES,
    CurationTypes,
)
from curation.constants.models.evidence import (
    TYPING_METHOD_CHOICES,
    ZYGOSITY_CHOICES,
    TypingMethod,
    Zygosity,
)
from curation.models import Curation, Demographic, Evidence
from disease.models import Disease
from haplotype.models import Haplotype
from publication.models import Publication
from repo.constants import PUBLIC_CURATION_FIELDS, PUBLIC_EVIDENCE_FIELDS
from repo.models import PublishedCuration
from repo.serializers import (
    CLASSIFICATION_COLUMNS,
    EVIDENCE_COLUMN_NAMES,
    EVIDENCE_COLUMNS,
    EXPORT_NAMES,
    flatten_classification_rows,
    flatten_evidence,
    flatten_published_curation,
    serialize_evidence,
    serialize_published_curation,
)

if TYPE_CHECKING:
    from django.test.client import _MonkeyPatchedWSGIResponse


class PublishedCurationModelTest(TestCase):
    fixtures = ["test_alleles.json", "test_diseases.json"]

    def setUp(self):
        self.allele = Allele.objects.get(pk=1)
        self.disease = Disease.objects.get(pk=1)
        self.user = User.objects.create_user(username="testuser", password="testpass")  # ruff: ignore[hardcoded-password-func-arg]
        self.curation = Curation.objects.create(
            curation_type=CurationTypes.ALLELE,
            allele=self.allele,
            disease=self.disease,
            status=Status.APPROVED,
        )

    def test_create_published_curation(self):
        published = PublishedCuration.objects.create(
            curation=self.curation,
            published_by=self.user,
        )
        self.assertEqual(published.curation, self.curation)
        self.assertEqual(published.published_by, self.user)
        self.assertEqual(published.version, 1)
        self.assertIsNotNone(published.published_at)

    def test_string_representation(self):
        self.curation.save()  # Ensure slug is generated.
        published = PublishedCuration.objects.create(
            curation=self.curation,
            published_by=self.user,
        )
        self.assertEqual(str(published), f"Published: {self.curation.slug}")

    def test_one_to_one_constraint(self):
        PublishedCuration.objects.create(
            curation=self.curation,
            published_by=self.user,
        )
        with self.assertRaises(IntegrityError):
            PublishedCuration.objects.create(
                curation=self.curation,
                published_by=self.user,
            )

    def test_reverse_relationship(self):
        published = PublishedCuration.objects.create(
            curation=self.curation,
            published_by=self.user,
        )
        self.assertEqual(self.curation.publication, published)  # type: ignore

    def test_get_absolute_url(self):
        self.curation.save()  # Ensure slug is generated.
        published = PublishedCuration.objects.create(
            curation=self.curation,
            published_by=self.user,
        )
        expected_url = reverse(
            "repo-detail", kwargs={"curation_slug": self.curation.slug}
        )
        self.assertEqual(published.get_absolute_url(), expected_url)


class CurationPublishViewTest(ProtectedViewTestMixin, TestCase):
    fixtures = ["test_alleles.json", "test_diseases.json"]
    # Publish is a redirect action, not a page view, so these don't apply.
    template = ""
    page_name = ""
    expected_text: list[str] = []

    def setUp(self):
        self.url = ""  # Set a placeholder; will be updated after super().setUp()
        super().setUp()
        self.allele = Allele.objects.get(pk=1)
        self.disease = Disease.objects.get(pk=1)
        self.curation = Curation.objects.create(
            curation_type=CurationTypes.ALLELE,
            allele=self.allele,
            disease=self.disease,
            status=Status.APPROVED,
        )
        self.curation.save()  # Ensure slug is generated.
        self.url = reverse(
            "curation-publish", kwargs={"curation_slug": self.curation.slug}
        )

    # Publish is a redirect action, not a page view, so skip inherited page tests.
    def test_template(self):
        pass

    def test_page_name_in_response(self):
        pass

    def test_expected_text_in_response(self):
        pass

    @override
    def test_permission_granted_if_yes_phi_yes_perms(self):
        # Publish always redirects (302) rather than returning 200, so we override.
        self.client.force_login(self.user4_yes_phi_yes_perms)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 302)

    def test_publish_approved_curation(self):
        self.client.force_login(self.user4_yes_phi_yes_perms)
        response = self.client.post(self.url)

        self.assertEqual(response.status_code, 302)
        self.assertEqual(PublishedCuration.objects.count(), 1)
        self.curation.refresh_from_db()
        self.assertEqual(self.curation.status, Status.PUBLISHED)

    def test_cannot_publish_in_progress_curation(self):
        curation = Curation.objects.create(
            curation_type=CurationTypes.ALLELE,
            allele=self.allele,
            disease=self.disease,
            status=Status.IN_PROGRESS,
        )
        curation.save()  # Ensure slug is generated.
        url = reverse("curation-publish", kwargs={"curation_slug": curation.slug})
        self.client.force_login(self.user4_yes_phi_yes_perms)
        response = self.client.post(url)

        self.assertEqual(response.status_code, 302)
        self.assertEqual(PublishedCuration.objects.count(), 0)

    def test_cannot_publish_already_published_curation(self):
        self.curation.status = Status.PUBLISHED
        self.curation.save()
        PublishedCuration.objects.create(
            curation=self.curation,
            published_by=self.user4_yes_phi_yes_perms,
        )
        self.client.force_login(self.user4_yes_phi_yes_perms)
        response = self.client.post(self.url)

        self.assertEqual(response.status_code, 302)
        self.assertEqual(PublishedCuration.objects.count(), 1)

    def test_get_does_not_publish(self):
        self.client.force_login(self.user4_yes_phi_yes_perms)
        response = self.client.get(self.url)

        self.assertEqual(response.status_code, 302)
        self.assertEqual(PublishedCuration.objects.count(), 0)
        self.curation.refresh_from_db()
        self.assertEqual(self.curation.status, Status.APPROVED)


class RepoSearchViewTest(TestCase):
    fixtures = ["test_alleles.json", "test_diseases.json"]

    def setUp(self):
        self.client = Client()
        self.url = reverse("repo-search")

    def test_response_code(self):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)

    def test_displays_published_curations(self):
        allele = Allele.objects.get(pk=1)
        disease = Disease.objects.get(pk=1)
        user = User.objects.create_user(username="testuser", password="testpass")  # ruff: ignore[hardcoded-password-func-arg]

        curation = Curation.objects.create(
            curation_type=CurationTypes.ALLELE,
            allele=allele,
            disease=disease,
            status=Status.PUBLISHED,
        )
        curation.save()
        PublishedCuration.objects.create(
            curation=curation,
            published_by=user,
        )

        response = self.client.get(self.url)
        self.assertContains(response, curation.slug)
        self.assertContains(response, "No Classification Set")

    def test_json_button_uses_common_icon_and_link_color(self):
        user = User.objects.create_user(username="testuser", password="testpass")  # ruff: ignore[hardcoded-password-func-arg]
        curation = Curation.objects.create(
            curation_type=CurationTypes.ALLELE,
            allele=Allele.objects.get(pk=1),
            disease=Disease.objects.get(pk=1),
            status=Status.PUBLISHED,
        )
        PublishedCuration.objects.create(curation=curation, published_by=user)

        response = self.client.get(self.url)
        url = reverse("repo-download-single", args=[curation.slug])
        self.assertContains(
            response, f'<a href="{url}" class="button is-small is-link is-light">'
        )
        self.assertContains(response, '<i class="bi bi-download mr-2"></i>')

    def _publish(self, allele_pk: int, published_at: datetime) -> Curation:
        user, _ = User.objects.get_or_create(username="testuser")
        curation = Curation.objects.create(
            curation_type=CurationTypes.ALLELE,
            allele=Allele.objects.get(pk=allele_pk),
            disease=Disease.objects.get(pk=1),
            status=Status.PUBLISHED,
        )
        published = PublishedCuration.objects.create(
            curation=curation, published_by=user
        )
        PublishedCuration.objects.filter(pk=published.pk).update(
            published_at=published_at
        )
        return curation

    def test_lists_most_recently_published_first(self):
        # Create the newer one first so updated_at order differs from publish order.
        newer = self._publish(2, timezone.now() - timedelta(days=1))
        older = self._publish(1, timezone.now() - timedelta(days=30))

        content = self.client.get(self.url).content.decode()
        self.assertLess(content.index(newer.slug), content.index(older.slug))

    def test_table_shows_published_date_and_no_type(self):
        published_at = timezone.now() - timedelta(days=30)
        self._publish(1, published_at)

        response = self.client.get(self.url)
        columns = response.context["table"].columns
        self.assertNotIn("curation_type", columns.names())
        self.assertEqual(columns["published_at"].header, "Published")
        self.assertContains(response, published_at.strftime("%Y-%m-%d"))


class PublishedCurationDetailViewTest(TestCase):
    fixtures = ["test_alleles.json", "test_diseases.json", "test_haplotypes.json"]

    def setUp(self):
        self.client = Client()
        allele = Allele.objects.get(pk=1)
        disease = Disease.objects.get(pk=1)
        user = User.objects.create_user(username="testuser", password="testpass")  # ruff: ignore[hardcoded-password-func-arg]

        self.curation = Curation.objects.create(
            curation_type=CurationTypes.ALLELE,
            allele=allele,
            disease=disease,
            status=Status.PUBLISHED,
        )
        self.curation.save()
        self.published = PublishedCuration.objects.create(
            curation=self.curation,
            published_by=user,
        )

    def test_shows_the_taskforce_name_and_id(self):
        Curation.objects.filter(pk=self.curation.pk).update(
            taskforce=HLA_CURATION_TASKFORCE_ID
        )
        response = self.client.get(reverse("repo-detail", args=[self.curation.slug]))
        name = dict(TASKFORCE_CHOICES)[HLA_CURATION_TASKFORCE_ID]
        self.assertContains(response, f"{name} ({HLA_CURATION_TASKFORCE_ID})")

    def test_response_code(self):
        url = reverse("repo-detail", kwargs={"curation_slug": self.curation.slug})
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)

    def test_displays_curation_details(self):
        url = reverse("repo-detail", kwargs={"curation_slug": self.curation.slug})
        response = self.client.get(url)
        self.assertContains(response, self.curation.slug)
        self.assertContains(response, "Download as JSON")

    def test_displays_evidence_summary(self):
        self.curation.taskforce_evidence_summary = (
            "The taskforce found strong evidence."
        )
        self.curation.save()
        url = reverse("repo-detail", kwargs={"curation_slug": self.curation.slug})
        response = self.client.get(url)
        self.assertContains(response, "Evidence Summary")
        self.assertContains(response, "The taskforce found strong evidence.")

    def test_copy_button_uses_the_confirm_modal(self):
        curator = User.objects.create_user(username="curator", password="pw")  # ruff: ignore[hardcoded-password-func-arg]
        UserProfile.objects.create(
            user=curator, has_signed_phi_agreement=True, has_curation_permissions=True
        )
        self.client.force_login(curator)
        url = reverse("repo-detail", kwargs={"curation_slug": self.curation.slug})
        response = self.client.get(url)
        self.assertContains(response, "Copy and Recurate")
        self.assertContains(response, "data-confirm=")
        self.assertNotContains(response, "confirm(")

    def test_displays_taskforce_review_date(self):
        self.curation.taskforce_review_date = date(2026, 9, 1)
        self.curation.save()
        url = reverse("repo-detail", kwargs={"curation_slug": self.curation.slug})
        response = self.client.get(url)
        self.assertContains(response, "Review Date")
        self.assertContains(response, "2026-09-01")

    def test_displays_override_reason_when_set(self):
        url = reverse("repo-detail", kwargs={"curation_slug": self.curation.slug})
        self.assertNotContains(self.client.get(url), "Override Reason")
        self.curation.taskforce_override_reason = "Replicated in two cohorts."
        self.curation.save()
        response = self.client.get(url)
        self.assertContains(response, "Override Reason")
        self.assertContains(response, "Replicated in two cohorts.")

    def test_displays_car_linkout_for_allele(self):
        url = reverse("repo-detail", kwargs={"curation_slug": self.curation.slug})
        response = self.client.get(url)
        self.assertContains(
            response, 'href="https://reg.clinicalgenome.org/allele/ui/hla/id/XAHLA123"'
        )

    def test_displays_mondo_linkout(self):
        url = reverse("repo-detail", kwargs={"curation_slug": self.curation.slug})
        response = self.client.get(url)
        self.assertContains(response, 'href="http://purl.obolibrary.org/obo/MONDO_123"')
        self.assertContains(response, "MONDO:123")

    def test_displays_car_linkout_for_each_haplotype_allele(self):
        haplotype = Haplotype.objects.get(pk=1)
        curation = Curation.objects.create(
            curation_type=CurationTypes.HAPLOTYPE,
            haplotype=haplotype,
            disease=Disease.objects.get(pk=1),
            status=Status.PUBLISHED,
        )
        PublishedCuration.objects.create(curation=curation)
        url = reverse("repo-detail", kwargs={"curation_slug": curation.slug})
        response = self.client.get(url)
        self.assertContains(response, haplotype.display_name)
        for car_id in ("XAHLA123", "XAHLA456"):
            self.assertContains(
                response,
                f'href="https://reg.clinicalgenome.org/allele/ui/hla/id/{car_id}"',
            )

    def test_displays_additional_notes(self):
        self.curation.taskforce_additional_notes = "Taskforce note."
        self.curation.save()
        url = reverse("repo-detail", kwargs={"curation_slug": self.curation.slug})
        response = self.client.get(url)
        self.assertContains(response, "Additional Notes")
        self.assertContains(response, "Taskforce note.")

    def test_linkout_has_no_space_inside_parentheses(self):
        url = reverse("repo-detail", kwargs={"curation_slug": self.curation.slug})
        response = self.client.get(url)
        self.assertContains(
            response, '(<a href="http://purl.obolibrary.org/obo/MONDO_123"'
        )
        self.assertContains(response, "</a>)")


class JSONDownloadViewTest(TestCase):
    fixtures = ["test_alleles.json", "test_diseases.json", "test_haplotypes.json"]

    def setUp(self):
        self.client = Client()
        allele = Allele.objects.get(pk=1)
        disease = Disease.objects.get(pk=1)
        user = User.objects.create_user(username="testuser", password="testpass")  # ruff: ignore[hardcoded-password-func-arg]

        self.curation = Curation.objects.create(
            curation_type=CurationTypes.ALLELE,
            allele=allele,
            disease=disease,
            status=Status.PUBLISHED,
        )
        self.curation.save()
        self.published = PublishedCuration.objects.create(
            curation=self.curation,
            published_by=user,
        )

    def test_download_single_json(self):
        url = reverse(
            "repo-download-single", kwargs={"curation_slug": self.curation.slug}
        )
        response = self.client.get(url)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/json")
        self.assertIn("attachment", response["Content-Disposition"])

        data = json.loads(response.content)
        self.assertEqual(data["curation"]["curation_id"], self.curation.slug)

    def test_json_includes_taskforce_review_date(self):
        url = reverse(
            "repo-download-single", kwargs={"curation_slug": self.curation.slug}
        )
        data = json.loads(self.client.get(url).content)
        self.assertIsNone(data["curation"]["curation"]["taskforce_review_date"])

        self.curation.taskforce_review_date = date(2026, 9, 1)
        self.curation.save()
        data = json.loads(self.client.get(url).content)
        self.assertEqual(
            data["curation"]["curation"]["taskforce_review_date"], "2026-09-01"
        )

    def test_json_includes_override_reason(self):
        self.curation.taskforce_override_reason = "Replicated in two cohorts."
        self.curation.save()
        url = reverse(
            "repo-download-single", kwargs={"curation_slug": self.curation.slug}
        )
        data = json.loads(self.client.get(url).content)
        self.assertEqual(
            data["curation"]["curation"]["taskforce_override_reason"],
            "Replicated in two cohorts.",
        )

    def test_json_has_no_expert_panel_keys(self):
        url = reverse(
            "repo-download-single", kwargs={"curation_slug": self.curation.slug}
        )
        keys = json.loads(self.client.get(url).content)["curation"]["curation"]
        self.assertEqual(
            [key for key in keys if key == "ep" or key.startswith("ep_")], []
        )

    def test_curation_keys_are_the_public_fields_plus_structure(self):
        data = serialize_published_curation(self.published)["curation"]
        public = {EXPORT_NAMES.get(name, name) for name in PUBLIC_CURATION_FIELDS}
        structural = {"status", "copied_from", "score", "type", "allele", "disease"}
        self.assertEqual(set(data) - {"added_at"}, public | structural)

    def test_evidence_keys_are_the_public_fields_plus_structure(self):
        publication = Publication.objects.create(
            slug="P999999", title="T", author="A", publication_year=2020, pubmed_id="1"
        )
        evidence = Evidence.objects.create(
            curation=self.curation, publication=publication
        )
        structural = {
            "evidence_id",
            "status",
            "is_included",
            "publication",
            "score",
            "added_at",
        }
        self.assertEqual(
            set(serialize_evidence(evidence)), PUBLIC_EVIDENCE_FIELDS | structural
        )

    def test_download_all_json(self):
        url = reverse("repo-download-all")
        response = self.client.get(url)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/json")
        self.assertIn("attachment", response["Content-Disposition"])

        data = json.loads(response.content)
        self.assertEqual(data["total_count"], 1)
        self.assertEqual(len(data["published_curations"]), 1)
        self.assertEqual(
            data["published_curations"][0]["curation_id"], self.curation.slug
        )

    def test_json_includes_disease_iri(self):
        url = reverse(
            "repo-download-single", kwargs={"curation_slug": self.curation.slug}
        )
        data = json.loads(self.client.get(url).content)
        self.assertEqual(
            data["curation"]["curation"]["disease"]["iri"],
            "http://purl.obolibrary.org/obo/MONDO_123",
        )

    def test_json_includes_haplotype_allele_car_id(self):
        curation = Curation.objects.create(
            curation_type=CurationTypes.HAPLOTYPE,
            haplotype=Haplotype.objects.get(pk=1),
            disease=Disease.objects.get(pk=1),
            status=Status.PUBLISHED,
        )
        PublishedCuration.objects.create(curation=curation)
        url = reverse("repo-download-single", kwargs={"curation_slug": curation.slug})
        data = json.loads(self.client.get(url).content)
        alleles = data["curation"]["curation"]["haplotype"]["alleles"]
        self.assertEqual(sorted(a["car_id"] for a in alleles), ["XAHLA123", "XAHLA456"])


class ReadOnlyEnforcementTest(TestCase):
    fixtures = ["test_alleles.json", "test_diseases.json"]

    def setUp(self):
        self.client = Client()
        allele = Allele.objects.get(pk=1)
        disease = Disease.objects.get(pk=1)
        self.user = User.objects.create_user(username="testuser", password="testpass")  # ruff: ignore[hardcoded-password-func-arg]
        UserProfile.objects.create(
            user=self.user,
            has_signed_phi_agreement=True,
            has_curation_permissions=True,
        )
        self.client.force_login(self.user)

        self.curation = Curation.objects.create(
            curation_type=CurationTypes.ALLELE,
            allele=allele,
            disease=disease,
            status=Status.PUBLISHED,
        )
        self.curation.save()
        self.published = PublishedCuration.objects.create(
            curation=self.curation,
            published_by=self.user,
        )

    def test_cannot_edit_published_evidence(self):
        url = reverse(
            "curation-edit-evidence", kwargs={"curation_slug": self.curation.slug}
        )
        response = self.client.get(url)

        self.assertEqual(response.status_code, 302)
        self.assertRedirects(
            response,
            reverse("curation-detail", kwargs={"curation_slug": self.curation.slug}),
        )


class SupersessionTest(TestCase):
    fixtures = ["test_alleles.json", "test_diseases.json"]

    def setUp(self):
        self.allele = Allele.objects.get(pk=1)
        self.disease = Disease.objects.get(pk=1)
        self.user = User.objects.create_user(username="superuser_t", password="pw")  # ruff: ignore[hardcoded-password-func-arg]
        UserProfile.objects.create(
            user=self.user,
            has_signed_phi_agreement=True,
            has_curation_permissions=True,
        )

    def _make_published(self, copied_from: Curation | None = None) -> PublishedCuration:
        curation = Curation.objects.create(
            curation_type=CurationTypes.ALLELE,
            allele=self.allele,
            disease=self.disease,
            status=Status.PUBLISHED,
            copied_from=copied_from,
        )
        return PublishedCuration.objects.create(
            curation=curation, published_by=self.user
        )

    def test_is_superseded_false_when_no_copy(self):
        from repo.views import is_superseded

        published = self._make_published()
        self.assertFalse(is_superseded(published))

    def test_is_superseded_true_when_copy_is_published(self):
        from repo.views import is_superseded

        original = self._make_published()
        self._make_published(copied_from=original.curation)
        self.assertTrue(is_superseded(original))

    def test_is_superseded_true_for_multi_hop_chain(self):
        from repo.views import is_superseded

        original = self._make_published()
        copy1 = self._make_published(copied_from=original.curation)
        self._make_published(copied_from=copy1.curation)
        self.assertTrue(is_superseded(original))

    def test_get_superseding_returns_none_when_not_superseded(self):
        from repo.views import get_superseding

        published = self._make_published()
        self.assertIsNone(get_superseding(published))

    def test_get_superseding_returns_copy(self):
        from repo.views import get_superseding

        original = self._make_published()
        copy_pub = self._make_published(copied_from=original.curation)
        result = get_superseding(original)
        self.assertEqual(result, copy_pub)


class CopyButtonTest(TestCase):
    fixtures = ["test_alleles.json", "test_diseases.json"]

    def setUp(self):
        self.client = Client()
        self.allele = Allele.objects.get(pk=1)
        self.disease = Disease.objects.get(pk=1)
        self.user = User.objects.create_user(username="copy_button_user", password="pw")  # ruff: ignore[hardcoded-password-func-arg]
        UserProfile.objects.create(
            user=self.user,
            has_signed_phi_agreement=True,
            has_curation_permissions=True,
        )
        self.curation = Curation.objects.create(
            curation_type=CurationTypes.ALLELE,
            allele=self.allele,
            disease=self.disease,
            status=Status.PUBLISHED,
        )
        self.curation.save()
        self.published = PublishedCuration.objects.create(
            curation=self.curation, published_by=self.user
        )

    def test_copy_button_visible_for_curators(self):
        self.client.force_login(self.user)
        url = reverse("repo-detail", kwargs={"curation_slug": self.curation.slug})
        response = self.client.get(url)
        self.assertContains(response, "Copy and Recurate")


def assert_no_login_links(
    test: TestCase, response: "_MonkeyPatchedWSGIResponse"
) -> None:
    """Asserts every internal link on a page loads for an anonymous user.

    Fetches each internal href with a fresh anonymous client and fails if one
    doesn't return 200, which catches links that send visitors to the login page.
    """
    soup = BeautifulSoup(response.content, "html.parser")
    hrefs = {str(a["href"]) for a in soup.find_all("a", href=True)}
    internal = {
        href
        for href in hrefs
        if href.startswith("/") and not href.startswith(settings.STATIC_URL)
    }
    test.assertTrue(internal, "Expected the page to have internal links.")
    for href in sorted(internal):
        with test.subTest(href=href):
            linked = Client().get(href)
            test.assertEqual(
                linked.status_code,
                200,
                f"{href} returned {linked.status_code}"
                f" (Location: {linked.get('Location')}).",
            )
            test.assertFalse(linked.get("Location", "").startswith(settings.LOGIN_URL))


class PublicLayoutTest(TestCase):
    fixtures = [
        "test_alleles.json",
        "test_diseases.json",
        "test_publications.json",
        "test_curations.json",
        "test_evidence.json",
    ]

    def setUp(self):
        self.curation = Curation.objects.get(slug="C000001")
        Curation.objects.filter(pk=self.curation.pk).update(status=Status.PUBLISHED)
        user = User.objects.create_user(username="publisher")
        self.published = PublishedCuration.objects.create(
            curation=self.curation, published_by=user
        )
        self.evidence = self.curation.evidence.get()  # type: ignore
        slug = self.curation.slug
        history_id = self.published.history.get().history_id  # type: ignore
        self.urls = [
            reverse("repo-search"),
            reverse("repo-detail", args=[slug]),
            reverse("repo-history", args=[slug]),
            reverse("repo-change", args=[slug, history_id]),
        ]

    def test_internal_links_work_without_login(self):
        for url in self.urls:
            with self.subTest(url=url):
                assert_no_login_links(self, self.client.get(url))

    def test_pages_use_the_public_layout(self):
        for url in self.urls:
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertTemplateUsed(response, "layouts/public.html")
                self.assertTemplateNotUsed(response, "layouts/base.html")
                self.assertTemplateNotUsed(response, "partials/navbar.html")

    def test_pages_have_no_curator_navigation(self):
        for url in self.urls:
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertNotContains(response, "Log In")
                self.assertNotContains(response, 'id="log-out-button"')
                self.assertNotContains(response, f'href="{reverse("allele-list")}"')
                self.assertNotContains(response, f'href="{reverse("about")}"')

    def test_pages_link_back_to_the_hci(self):
        for url in self.urls:
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertContains(response, 'id="hci-link"')
                self.assertContains(response, f'href="{reverse("home")}"')

    def test_navbar_has_no_hlarepo_link(self):
        response = self.client.get(self.urls[1])
        navbar = response.content.decode().split('<nav class="navbar"')[1]
        navbar = navbar.split("</nav>")[0]
        self.assertNotIn(f'href="{reverse("repo-search")}"', navbar)
        self.assertIn('id="hci-link"', navbar)

    def test_breadcrumbs_start_at_hlarepo_home(self):
        for url in self.urls:
            with self.subTest(url=url):
                response = self.client.get(url)
                breadcrumb = response.content.decode().split('class="breadcrumb"')[1]
                breadcrumb = breadcrumb.split("</nav>")[0]
                self.assertIn("bi-database", breadcrumb)
                self.assertIn("HLArepo Home", breadcrumb)
                self.assertNotIn("bi-house", breadcrumb)

    def test_footer_names_hlarepo(self):
        response = self.client.get(self.urls[0])
        self.assertTemplateUsed(response, "partials/public_footer.html")
        self.assertTemplateUsed(response, "partials/footer_logos.html")
        self.assertContains(response, "The HLArepo is an open-source project.")
        self.assertNotContains(response, "The HCI is an open-source project.")

    def test_title_names_hlarepo(self):
        response = self.client.get(self.urls[1])
        self.assertContains(response, "<title>HLArepo |")

    def test_no_activation_notice_for_logged_in_non_curator(self):
        user = User.objects.create_user(username="visitor")
        UserProfile.objects.create(
            user=user, has_signed_phi_agreement=False, has_curation_permissions=False
        )
        self.client.force_login(user)
        response = self.client.get(self.urls[1])
        self.assertTemplateNotUsed(response, "partials/account_activation.html")

    def test_evidence_ids_are_plain_text(self):
        response = self.client.get(self.urls[1])
        self.assertContains(response, self.evidence.slug)
        evidence_url = reverse(
            "evidence-detail",
            kwargs={
                "curation_slug": self.curation.slug,
                "evidence_slug": self.evidence.slug,
            },
        )
        self.assertNotContains(response, evidence_url)
        self.assertNotContains(response, "Edit Evidence")

    def test_evidence_table_hides_needs_review(self):
        self.assertNotIn("needs_review", PUBLIC_EVIDENCE_FIELDS)
        response = self.client.get(self.urls[1])
        self.assertNotContains(response, "Needs Review")


class RepoAboutViewTest(OpenViewTestMixin, TestCase):
    url = reverse("repo-about")
    template = "repo/about.html"
    page_name = "About"
    expected_text = ["HLArepo", "published", "HLA Curation Interface", "Stanford"]


class RepoContactViewTest(OpenViewTestMixin, TestCase):
    url = reverse("repo-contact")
    template = "repo/contact.html"
    page_name = "Contact"
    expected_text = ["HLArepo", "hci@clinicalgenome.org"]


class RepoHelpViewTest(OpenViewTestMixin, TestCase):
    url = reverse("repo-help")
    template = "repo/help.html"
    page_name = "Help"
    expected_text = ["HLArepo", "hci@clinicalgenome.org", "subject"]


class RepoCitingViewTest(OpenViewTestMixin, TestCase):
    url = reverse("repo-citing")
    template = "repo/citing.html"
    page_name = "Citing"
    expected_text = [
        "Clinical Genome Resource. HLArepo. Stanford University.",
        "https://hci.clinicalgenome.org/repo/",
        "downloaded",
    ]


class RepoAcknowledgementsViewTest(OpenViewTestMixin, TestCase):
    url = reverse("repo-acknowledgements")
    template = "repo/acknowledgements.html"
    page_name = "Acknowledgements"
    expected_text = ["HLArepo", "U24HG009649", "Steven Mack"]


class RepoCollaboratorsViewTest(OpenViewTestMixin, TestCase):
    url = reverse("repo-collaborators")
    template = "repo/collaborators.html"
    page_name = "Collaborators"
    expected_text = ["The Baylor College of Medicine ClinGen Team", "ClinPGx"]


REPO_INFO_URL_NAMES = [
    "repo-about",
    "repo-contact",
    "repo-citing",
    "repo-help",
    "repo-acknowledgements",
    "repo-collaborators",
]
CORE_INFO_URL_NAMES = [
    "about",
    "contact",
    "citing",
    "help",
    "acknowledgements",
    "collaborators",
]


class RepoInfoPagesTest(TestCase):
    def test_pages_use_the_public_layout_without_login_links(self):
        for name in REPO_INFO_URL_NAMES:
            with self.subTest(name=name):
                response = self.client.get(reverse(name))
                self.assertTemplateUsed(response, "layouts/public.html")
                self.assertContains(response, "HLArepo Home")
                assert_no_login_links(self, response)

    def test_public_footer_links_to_repo_info_pages(self):
        response = self.client.get(reverse("repo-search"))
        for name in REPO_INFO_URL_NAMES:
            with self.subTest(name=name):
                self.assertContains(response, f'href="{reverse(name)}"')
        for name in CORE_INFO_URL_NAMES:
            with self.subTest(name=name):
                self.assertNotContains(response, f'href="{reverse(name)}"')


class RepoLandingPageTest(TestCase):
    url = reverse("repo-search")

    def test_has_title_subtitle_and_description(self):
        response = self.client.get(self.url)
        self.assertContains(response, '<h1 class="title">HLArepo</h1>', html=True)
        self.assertContains(response, 'class="subtitle"')
        self.assertContains(response, 'id="repo-description"')
        self.assertContains(response, "HLA Curation Interface")

    def test_shows_the_hci_logo(self):
        response = self.client.get(self.url)
        self.assertContains(response, f'src="{static("hci/img/hci-logo.png")}"')

    def test_downloads_section_has_the_json_download(self):
        response = self.client.get(self.url)
        soup = BeautifulSoup(response.content, "html.parser")
        downloads = soup.find(id="downloads")
        self.assertIsNotNone(downloads)
        assert downloads is not None
        link = downloads.find(id="download-button")
        assert link is not None
        self.assertEqual(link["href"], reverse("repo-download-all"))

    def test_htmx_search_returns_only_results(self):
        response = self.client.get(self.url, {"q": "x"}, headers={"HX-Request": "true"})
        self.assertTemplateUsed(response, "common/partials/search_results.html")
        self.assertTemplateNotUsed(response, "layouts/public.html")
        self.assertNotContains(response, 'id="downloads"')


class CSVDownloadViewTest(TestCase):
    fixtures = [
        "test_alleles.json",
        "test_diseases.json",
        "test_haplotypes.json",
        "demographics.json",
    ]
    url = reverse("repo-download-all-csv")

    def setUp(self):
        self.user = User.objects.create_user(username="publisher")
        self.allele_curation = Curation.objects.create(
            curation_type=CurationTypes.ALLELE,
            allele=Allele.objects.get(pk=1),
            disease=Disease.objects.get(pk=1),
            status=Status.PUBLISHED,
            taskforce_evidence_summary='Line one, with a comma.\nLine "two".',
        )
        self.haplotype_curation = Curation.objects.create(
            curation_type=CurationTypes.HAPLOTYPE,
            haplotype=Haplotype.objects.get(pk=1),
            disease=Disease.objects.get(pk=1),
            status=Status.PUBLISHED,
        )
        self.publication = Publication.objects.create(
            slug="P999999", title="T", author="A", publication_year=2020, pubmed_id="1"
        )
        self.first_evidence = Evidence.objects.create(
            curation=self.allele_curation,
            publication=self.publication,
            zygosity=Zygosity.MONOALLELIC,
            is_gwas=True,
        )
        self.second_evidence = Evidence.objects.create(
            curation=self.allele_curation, publication=self.publication
        )
        # Publish the allele curation more recently, but create it first.
        self._publish(self.allele_curation, timezone.now() - timedelta(days=1))
        self._publish(self.haplotype_curation, timezone.now() - timedelta(days=30))

    def _publish(self, curation: Curation, published_at: datetime) -> None:
        published = PublishedCuration.objects.create(
            curation=curation, published_by=self.user
        )
        PublishedCuration.objects.filter(pk=published.pk).update(
            published_at=published_at
        )

    def _rows(self) -> list[dict[str, str]]:
        response = self.client.get(self.url)
        text = response.content.decode("utf-8-sig")
        return list(csv.DictReader(io.StringIO(text)))

    def test_response_is_a_csv_attachment(self):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response["Content-Type"].startswith("text/csv"))
        self.assertRegex(
            response["Content-Disposition"],
            r'^attachment; filename="hla_curations_all_\d{4}-\d{2}-\d{2}\.csv"$',
        )
        self.assertTrue(response.content.startswith(codecs.BOM_UTF8))

    def test_one_row_per_evidence_newest_classification_first(self):
        rows = self._rows()
        self.assertEqual(
            [(row["Curation ID"], row["Evidence ID"]) for row in rows],
            [
                (self.allele_curation.slug, self.first_evidence.slug),
                (self.allele_curation.slug, self.second_evidence.slug),
                (self.haplotype_curation.slug, ""),
            ],
        )

    def test_classification_without_evidence_has_blank_evidence_columns(self):
        haplotype_row = self._rows()[-1]
        for column in EVIDENCE_COLUMN_NAMES:
            with self.subTest(column=column):
                self.assertEqual(haplotype_row[column], "")

    def test_header_comes_from_the_row_builder(self):
        published = PublishedCuration.objects.get(curation=self.allele_curation)
        expected = list(flatten_classification_rows(published, url="")[0])
        response = self.client.get(self.url)
        header = next(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
        self.assertEqual(header, expected)

    def test_allele_and_haplotype_columns(self):
        allele_row, _, haplotype_row = self._rows()
        allele = self.allele_curation.allele
        assert allele is not None
        self.assertEqual(allele_row["Allele"], allele.name)
        self.assertEqual(allele_row["Allele CAR ID"], allele.car_id or "")
        self.assertEqual(allele_row["Haplotype"], "")
        haplotype = self.haplotype_curation.haplotype
        assert haplotype is not None
        self.assertEqual(haplotype_row["Haplotype"], haplotype.name)
        self.assertEqual(haplotype_row["Allele"], "")
        self.assertEqual(haplotype_row["Allele CAR ID"], "")

    def test_evidence_summary_round_trips(self):
        allele_row = self._rows()[0]
        self.assertEqual(
            allele_row["Taskforce Evidence Summary"],
            'Line one, with a comma.\nLine "two".',
        )

    def test_url_is_the_absolute_repo_url(self):
        allele_row = self._rows()[0]
        path = reverse("repo-detail", args=[self.allele_curation.slug])
        self.assertEqual(allele_row["URL"], f"http://testserver{path}")

    def test_evidence_values_are_readable(self):
        row = self._rows()[0]
        self.assertEqual(row["Zygosity"], ZYGOSITY_CHOICES[Zygosity.MONOALLELIC])
        self.assertEqual(row["GWAS"], "Yes")
        self.assertEqual(row["Publication Title"], self.publication.title)
        self.assertEqual(row["PubMed ID"], self.publication.pubmed_id)

    def test_has_only_public_columns(self):
        header = set(self._rows()[0].keys())
        for column in [*CLASSIFICATION_COLUMNS.values(), *EVIDENCE_COLUMNS.values()]:
            with self.subTest(column=column):
                self.assertIn(column, header)
        for column in header:
            with self.subTest(column=column):
                self.assertNotIn("Review", column.replace("Review Date", ""))
                self.assertNotIn("Notes", column.replace("Additional Notes", ""))
        self.assertNotIn("Published By", header)
        self.assertNotIn("Status", header)

    def test_query_count_does_not_grow_with_rows(self):
        Evidence.objects.create(
            curation=self.haplotype_curation, publication=self.publication
        )
        # One query for the curations and their related rows, one for evidence
        # and its publication, and one for demographics.
        with self.assertNumQueries(3):
            self.client.get(self.url)

    def test_landing_page_has_csv_button(self):
        response = self.client.get(reverse("repo-search"))
        soup = BeautifulSoup(response.content, "html.parser")
        downloads = soup.find(id="downloads")
        assert downloads is not None
        self.assertIsNotNone(downloads.find("a", href=self.url))


class XLSXDownloadViewTest(TestCase):
    fixtures = [
        "test_alleles.json",
        "test_diseases.json",
        "test_haplotypes.json",
        "demographics.json",
    ]
    url = reverse("repo-download-all-xlsx")

    def setUp(self):
        user = User.objects.create_user(username="publisher")
        self.curation = Curation.objects.create(
            curation_type=CurationTypes.ALLELE,
            allele=Allele.objects.get(pk=1),
            disease=Disease.objects.get(pk=1),
            status=Status.PUBLISHED,
            taskforce_evidence_summary="Summary.",
        )
        self.publication = Publication.objects.create(
            slug="P999999", title="T", author="A", publication_year=2020, pubmed_id="1"
        )
        self.evidence = Evidence.objects.create(
            curation=self.curation,
            publication=self.publication,
            zygosity=Zygosity.BIALLELIC,
            p_value=Decimal("0.0001"),
            is_included=True,
        )
        self.evidence.demographics.set(Demographic.objects.all()[:2])
        self.published = PublishedCuration.objects.create(
            curation=self.curation, published_by=user
        )

    def _workbook(self) -> Workbook:
        response = self.client.get(self.url)
        return load_workbook(io.BytesIO(response.content))

    @staticmethod
    def _sheet_rows(workbook: Workbook, name: str) -> list[dict[str, object]]:
        rows = list(workbook[name].iter_rows(values_only=True))
        header = [str(cell) for cell in rows[0]]
        return [
            {
                key: ("" if value is None else value)
                for key, value in zip(header, row, strict=True)
            }
            for row in rows[1:]
        ]

    def test_response_is_an_xlsx_attachment(self):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertRegex(
            response["Content-Disposition"],
            r'^attachment; filename="hla_curations_all_\d{4}-\d{2}-\d{2}\.xlsx"$',
        )

    def test_has_curations_and_evidence_sheets(self):
        self.assertEqual(self._workbook().sheetnames, ["Curations", "Evidence"])

    def test_curations_sheet_matches_the_flat_row(self):
        def normalize(value: object) -> object:
            # Excel stores 0.0 as 0, so compare numbers as numbers.
            if isinstance(value, int | float) and not isinstance(value, bool):
                return float(value)
            return str(value)

        url = f"http://testserver{self.published.get_absolute_url()}"
        expected = flatten_published_curation(self.published, url)
        (row,) = self._sheet_rows(self._workbook(), "Curations")
        self.assertEqual(
            {key: normalize(value) for key, value in row.items()},
            {key: normalize(value) for key, value in expected.items()},
        )

    def test_evidence_sheet_has_one_row_per_evidence(self):
        (row,) = self._sheet_rows(self._workbook(), "Evidence")
        self.assertEqual(list(row), ["Curation ID", *EVIDENCE_COLUMN_NAMES])
        self.assertEqual(row["Curation ID"], self.curation.slug)
        self.assertEqual(row["Evidence ID"], self.evidence.slug)
        self.assertEqual(row["Publication Title"], self.publication.title)
        self.assertEqual(str(row["PubMed ID"]), self.publication.pubmed_id)
        groups = [demo.group for demo in self.evidence.demographics.all()]
        self.assertEqual(row["Demographics"], "; ".join(groups))

    def test_values_are_readable_and_decimals_are_numbers(self):
        (row,) = self._sheet_rows(self._workbook(), "Evidence")
        self.assertEqual(row["Zygosity"], ZYGOSITY_CHOICES[Zygosity.BIALLELIC])
        self.assertEqual(row["Included in Score"], "Yes")
        p_value = row["p-value"]
        assert isinstance(p_value, float)
        self.assertAlmostEqual(p_value, 0.0001)

    def test_header_rows_are_bold_and_frozen(self):
        workbook = self._workbook()
        for name in workbook.sheetnames:
            with self.subTest(sheet=name):
                sheet = workbook[name]
                self.assertTrue(sheet["A1"].font.bold)
                self.assertEqual(sheet.freeze_panes, "A2")

    def test_query_count_does_not_grow_with_rows(self):
        other = Curation.objects.create(
            curation_type=CurationTypes.ALLELE,
            allele=Allele.objects.get(pk=2),
            disease=Disease.objects.get(pk=1),
            status=Status.PUBLISHED,
        )
        Evidence.objects.create(curation=other, publication=self.publication)
        PublishedCuration.objects.create(curation=other)
        # Curations, evidence, and demographics.
        with self.assertNumQueries(3):
            self.client.get(self.url)

    def test_landing_page_has_xlsx_button(self):
        response = self.client.get(reverse("repo-search"))
        soup = BeautifulSoup(response.content, "html.parser")
        downloads = soup.find(id="downloads")
        assert downloads is not None
        self.assertIsNotNone(downloads.find("a", href=self.url))


class FlattenPublishedCurationTest(TestCase):
    fixtures = ["test_alleles.json", "test_diseases.json", "demographics.json"]

    def setUp(self):
        self.curation = Curation.objects.create(
            curation_type=CurationTypes.ALLELE,
            allele=Allele.objects.get(pk=1),
            disease=Disease.objects.get(pk=1),
            status=Status.PUBLISHED,
            taskforce_review_date=date(2026, 9, 1),
        )
        self.published = PublishedCuration.objects.create(curation=self.curation)

    def test_starts_with_the_identifying_columns(self):
        row = flatten_published_curation(self.published, url="https://example.org")
        self.assertEqual(
            list(row)[:4], ["Curation ID", "URL", "Version", "Published (UTC)"]
        )
        self.assertEqual(row["URL"], "https://example.org")

    def test_formats_dates_and_blanks_missing_values(self):
        row = flatten_published_curation(self.published, url="")
        self.assertEqual(row["Taskforce Review Date"], "2026-09-01")
        self.assertEqual(row["Taskforce Override Reason"], "")
        self.assertEqual(row["Update Of"], "")
        self.assertNotIn(None, row.values())

    def test_classification_matches_the_page(self):
        row = flatten_published_curation(self.published, url="")
        self.assertEqual(row["Classification"], self.curation.classification_display)

    def test_taskforce_is_its_name_and_id(self):
        Curation.objects.filter(pk=self.curation.pk).update(
            taskforce=HLA_CURATION_TASKFORCE_ID
        )
        self.published.refresh_from_db()
        row = flatten_published_curation(self.published, url="")
        name = dict(TASKFORCE_CHOICES)[HLA_CURATION_TASKFORCE_ID]
        self.assertEqual(row["Taskforce"], f"{name} ({HLA_CURATION_TASKFORCE_ID})")

    def test_every_public_field_has_a_column(self):
        self.assertEqual(
            set(CLASSIFICATION_COLUMNS),
            PUBLIC_CURATION_FIELDS - {"taskforce_classification"},
        )
        self.assertEqual(set(EVIDENCE_COLUMNS), PUBLIC_EVIDENCE_FIELDS)


class FlattenEvidenceTest(TestCase):
    fixtures = ["test_alleles.json", "test_diseases.json", "demographics.json"]

    def setUp(self):
        curation = Curation.objects.create(
            curation_type=CurationTypes.ALLELE,
            allele=Allele.objects.get(pk=1),
            disease=Disease.objects.get(pk=1),
        )
        self.evidence = Evidence.objects.create(
            curation=curation,
            typing_method=TypingMethod.SEROLOGICAL,
            is_protective=False,
            odds_ratio=Decimal("2.5"),
        )

    def test_columns_match_the_blank_evidence_columns(self):
        self.assertEqual(list(flatten_evidence(self.evidence)), EVIDENCE_COLUMN_NAMES)

    def test_uses_display_labels(self):
        row = flatten_evidence(self.evidence)
        self.assertEqual(
            row["Typing Method"], TYPING_METHOD_CHOICES[TypingMethod.SEROLOGICAL]
        )
        self.assertEqual(row["Protective"], "No")

    def test_decimals_stay_numbers(self):
        row = flatten_evidence(self.evidence)
        self.assertEqual(row["Odds Ratio"], Decimal("2.5"))

    def test_joins_demographic_groups(self):
        groups = list(Demographic.objects.order_by("group")[:2])
        self.evidence.demographics.set(groups)
        row = flatten_evidence(self.evidence)
        for group in groups:
            with self.subTest(group=group.group):
                self.assertIn(group.group, str(row["Demographics"]))
        self.assertIn("; ", str(row["Demographics"]))
